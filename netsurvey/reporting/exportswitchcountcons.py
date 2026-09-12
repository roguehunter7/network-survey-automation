import logging
import math
import os
import sqlite3

import pandas as pd
from openpyxl.utils import get_column_letter

# --- Configuration ---
DB_FILENAME = "network_survey.db"
OUTPUT_FILENAME = "Consolidated_Proposed_Switches_Report.xlsx"

# --- Logging Setup ---
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)

# --- SQL Query Definition ---
AGGREGATED_QUERY = """
WITH all_switch_like_devices AS (
    -- Combine managed switches, excluding specific hostnames and those without port counts
    SELECT
        s.site, s.building, s.floor, s.lab_area_name, s.user_verified_used_ports AS used_ports
    FROM switches s
    WHERE
        s.user_verified_total_ports IS NOT NULL
        AND LOWER(COALESCE(s.hostname, '')) NOT LIKE '%swp%'
        AND LOWER(COALESCE(s.hostname, '')) NOT LIKE '%rss%'
    UNION ALL
    -- Combine other logged devices that have port counts
    SELECT
        o.site, o.building, o.floor, o.lab_area_name, o.reported_used_ports AS used_ports
    FROM logged_other_devices o
    WHERE o.reported_total_ports IS NOT NULL
)
-- Aggregate the data by lab area to calculate total used ports
SELECT
    site,
    building,
    floor,
    lab_area_name,
    SUM(COALESCE(used_ports, 0)) AS total_used_ports
FROM all_switch_like_devices
-- Ensure that the location identifiers are not null
WHERE site IS NOT NULL AND building IS NOT NULL AND floor IS NOT NULL AND lab_area_name IS NOT NULL
GROUP BY site, building, floor, lab_area_name
ORDER BY site, building, floor, lab_area_name;
"""


def get_db_connection(db_file):
    """Establishes a read-only connection to the SQLite database."""
    try:
        conn = sqlite3.connect(f"file:{db_file}?mode=ro", uri=True)
        logging.info(f"Successfully connected to database '{db_file}'.")
        return conn
    except sqlite3.Error as e:
        logging.error(f"Error connecting to database '{db_file}': {e}")
        return None


def fetch_aggregated_data(cursor):
    """Fetches and processes the aggregated switch data for all lab areas."""
    logging.info("Fetching aggregated data for all lab areas...")
    try:
        cursor.execute(AGGREGATED_QUERY)
        results = cursor.fetchall()
        logging.info(f"Found data for {len(results)} distinct lab areas.")
        return results
    except sqlite3.Error as e:
        logging.error(f"Failed to execute aggregated query: {e}")
        return []


def auto_adjust_column_width(worksheet):
    """Adjusts column width based on the maximum content length in each column."""
    for col_idx in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        for cell in worksheet[column_letter]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass  # Ignore any problematic cells
        adjusted_width = max_length + 3
        worksheet.column_dimensions[column_letter].width = adjusted_width


def create_consolidated_excel_report(report_data):
    """Creates a single Excel report with the proposed switch counts for all lab areas."""
    if not report_data:
        logging.warning("No data available to generate a report.")
        return

    report_rows = []
    for site, building, floor, lab_area, total_used_ports in report_data:
        # Calculate the proposed number of switches based on used ports
        proposed_24_port = math.ceil(total_used_ports / 24.0)
        proposed_48_port = math.ceil(total_used_ports / 48.0)
        report_rows.append(
            {
                "Site": site,
                "Building": building,
                "Floor": floor,
                "Lab Area": lab_area,
                "Proposed 24-Port Switches": proposed_24_port,
                "Proposed 48-Port Switches": proposed_48_port,
            }
        )

    df = pd.DataFrame(report_rows)

    try:
        with pd.ExcelWriter(OUTPUT_FILENAME, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Proposed Switch Counts", index=False)
            worksheet = writer.sheets["Proposed Switch Counts"]
            auto_adjust_column_width(worksheet)
        logging.info(
            f"Successfully generated consolidated Excel report: '{OUTPUT_FILENAME}'"
        )
    except Exception as e:
        logging.error(f"Failed to write Excel file: {e}", exc_info=True)


def main():
    """Main function to generate the consolidated report."""
    if not os.path.exists(DB_FILENAME):
        logging.error(
            f"Database file '{DB_FILENAME}' not found. Please ensure it is in the same directory."
        )
        return

    conn = get_db_connection(DB_FILENAME)
    if not conn:
        return

    try:
        cursor = conn.cursor()
        aggregated_data = fetch_aggregated_data(cursor)

        if not aggregated_data:
            logging.warning(
                "No data was retrieved from the database. The report cannot be generated."
            )
            return

        create_consolidated_excel_report(aggregated_data)

    except Exception as e:
        logging.error(
            f"An unexpected error occurred during report generation: {e}", exc_info=True
        )
    finally:
        if conn:
            conn.close()
            logging.info("Database connection closed.")
        print("\nExiting.")


if __name__ == "__main__":
    main()
