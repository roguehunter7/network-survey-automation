# exportcjexcel.py
import logging
import math
import os
import re
import sqlite3
import sys

import pandas as pd
from openpyxl.utils import get_column_letter

# --- Configuration ---
DB_FILENAME = "network_survey.db"

# --- Logging Setup ---
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)

# --- SQL Query Definitions ---
BASE_CTE = """
WITH all_switch_like_devices AS (
    SELECT
        s.serial_number, s.user_verified_total_ports AS total_ports, s.user_verified_used_ports AS used_ports,
        s.site, s.building, s.floor, s.lab_area_name,
        CASE WHEN EXISTS (SELECT 1 FROM switch_moonid_map m WHERE m.switch_serial_number = s.serial_number)
             THEN 1 ELSE 0 END AS is_in_moon
    FROM switches s
    WHERE
        s.user_verified_total_ports IS NOT NULL
        AND LOWER(COALESCE(s.hostname, '')) NOT LIKE '%swp%'
        AND LOWER(COALESCE(s.hostname, '')) NOT LIKE '%rss%'
    UNION ALL
    SELECT
        'lod-' || o.id AS serial_number, o.reported_total_ports AS total_ports, o.reported_used_ports AS used_ports,
        o.site, o.building, o.floor, o.lab_area_name,
        CASE WHEN o.reported_moonid IS NOT NULL AND o.reported_moonid != '' THEN 1 ELSE 0 END AS is_in_moon
    FROM logged_other_devices o
    WHERE o.reported_total_ports IS NOT NULL
)
"""

def get_queries_for_granularity(granularity="lab"):
    if granularity == "floor":
        grouping_where_clause = " WHERE site = ? AND building = ? AND floor = ?"
    else:  # 'lab'
        grouping_where_clause = " WHERE site = ? AND building = ? AND floor = ? AND lab_area_name = ?"

    queries = {
        "ge16_ports": BASE_CTE + f"SELECT COUNT(serial_number), SUM(COALESCE(used_ports, 0)), SUM(total_ports - COALESCE(used_ports, 0)) FROM all_switch_like_devices{grouping_where_clause} AND total_ports >= 16;",
        "in_moon": BASE_CTE + f"SELECT COUNT(serial_number), SUM(COALESCE(used_ports, 0)), SUM(total_ports - COALESCE(used_ports, 0)) FROM all_switch_like_devices{grouping_where_clause} AND is_in_moon = 1;",
        "not_in_moon": BASE_CTE + f"SELECT COUNT(serial_number), SUM(COALESCE(used_ports, 0)), SUM(total_ports - COALESCE(used_ports, 0)) FROM all_switch_like_devices{grouping_where_clause} AND is_in_moon = 0;",
        "lt16_ports": BASE_CTE + f"SELECT COUNT(serial_number), SUM(COALESCE(used_ports, 0)), SUM(total_ports - COALESCE(used_ports, 0)) FROM all_switch_like_devices{grouping_where_clause} AND total_ports < 16;",
    }
    return queries

BASE_AUDIT_LOG_SQL = """
SELECT
    site, building, floor, lab_area_name,
    identifier, serial_number, make, model, source_table, status, exclusion_reason
FROM (
    SELECT
        s.site, s.building, s.floor, s.lab_area_name, s.hostname AS identifier, s.serial_number, s.make, s.model,
        'Managed Switch' as source_table,
        CASE
            WHEN LOWER(COALESCE(s.hostname, '')) LIKE '%swp%' THEN 'Excluded'
            WHEN LOWER(COALESCE(s.hostname, '')) LIKE '%rss%' THEN 'Excluded'
            WHEN s.user_verified_total_ports IS NULL THEN 'Excluded'
            ELSE 'Included'
        END AS status,
        CASE
            WHEN LOWER(COALESCE(s.hostname, '')) LIKE '%swp%' THEN 'Hostname contains "swp"'
            WHEN LOWER(COALESCE(s.hostname, '')) LIKE '%rss%' THEN 'Hostname contains "rss"'
            WHEN s.user_verified_total_ports IS NULL THEN 'Port count is NULL'
            ELSE ''
        END AS exclusion_reason
    FROM switches s
    UNION ALL
    SELECT
        o.site, o.building, o.floor, o.lab_area_name, o.reported_serial_number AS identifier, o.reported_serial_number AS serial_number,
        o.reported_make AS make, o.reported_model AS model, 'Other Logged Device' as source_table,
        CASE WHEN o.reported_total_ports IS NULL THEN 'Excluded' ELSE 'Included' END AS status,
        CASE WHEN o.reported_total_ports IS NULL THEN 'Port count is NULL' ELSE '' END AS exclusion_reason
    FROM logged_other_devices o
)
"""

def get_db_connection(db_file):
    try:
        conn = sqlite3.connect(f"file:{db_file}?mode=ro", uri=True) # Read-only mode
        logging.info(f"Successfully connected to database '{db_file}'.")
        return conn
    except sqlite3.Error as e:
        logging.error(f"Error connecting to database '{db_file}': {e}")
        return None

def _prompt_user_for_choice(prompt_text, choices):
    if not choices:
        logging.warning(f"No choices available for: {prompt_text}")
        return None
    print(f"\n--- {prompt_text} ---")
    for i, choice in enumerate(choices, 1):
        print(f"  [{i}] {choice}")
    print("  [q] Quit")
    while True:
        user_input = input("Your choice: ").strip().lower()
        if user_input == 'q':
            return None
        try:
            choice_index = int(user_input) - 1
            if 0 <= choice_index < len(choices):
                return choices[choice_index]
            else:
                print("Invalid number. Please try again.")
        except ValueError:
            print("Invalid input. Please enter a number or 'q'.")

def fetch_distinct(cursor, query, params=()):
    cursor.execute(query, params)
    return sorted([row[0] for row in cursor.fetchall()])

def generate_report_for_scope(cursor, area_key, queries):
    ge16_count, ge16_used, ge16_unused = fetch_one_and_clean(cursor, queries["ge16_ports"], area_key)
    moon_count, moon_used, moon_unused = fetch_one_and_clean(cursor, queries["in_moon"], area_key)
    no_moon_count, no_moon_used, no_moon_unused = fetch_one_and_clean(cursor, queries["not_in_moon"], area_key)
    lt16_count, lt16_used, lt16_unused = fetch_one_and_clean(cursor, queries["lt16_ports"], area_key)
    categories_data = [
        ("Total Switches with >=16 ports", ge16_count, ge16_used, ge16_unused),
        ("Total Switches connected to known MOON", moon_count, moon_used, moon_unused),
        ("Total Switches Not in MOON", no_moon_count, no_moon_used, no_moon_unused),
        ("Total No.of Switches <16 ports", lt16_count, lt16_used, lt16_unused),
    ]
    report_rows = []
    for label, total, used, unused in categories_data:
        prop_24 = int(math.ceil(used / 24.0))
        prop_48 = int(math.ceil(used / 48.0))
        report_rows.append({
            "Category": label, "Total Switches": total, "Total Used Ports": used,
            "Total Unused Ports": unused, "Proposed 24 Port": prop_24,
            "Proposed 48 Port": prop_48, "Remarks": ""
        })
    return report_rows

def generate_scoped_audit_log(cursor, granularity, area_key):
    if granularity == "floor":
        where_clause = "WHERE site = ? AND building = ? AND floor = ?"
    else: # 'lab'
        where_clause = "WHERE site = ? AND building = ? AND floor = ? AND lab_area_name = ?"
    scoped_sql = f"{BASE_AUDIT_LOG_SQL} {where_clause} ORDER BY site, building, floor, lab_area_name, status, identifier;"
    try:
        logging.info("Generating data for the scoped Device Audit Log sheet...")
        df_audit = pd.read_sql_query(scoped_sql, cursor.connection, params=area_key)
        logging.info(f"Scoped audit log generated with {len(df_audit)} total entries.")
        return df_audit
    except Exception as e:
        logging.error(f"Failed to generate scoped audit log data: {e}")
        return pd.DataFrame()

def fetch_one_and_clean(cursor, query, params):
    try:
        cursor.execute(query, params)
        result = cursor.fetchone()
        return (tuple(x or 0 for x in result) if result and result[0] is not None else (0, 0, 0))
    except sqlite3.Error as e:
        logging.error(f"SQL Error during fetch: {e}\nQuery:\n{query[:500]}...")
        return (0, 0, 0)

# ==============================================================================
# ### THIS IS THE FIXED FUNCTION ###
# ==============================================================================
def auto_adjust_column_width(worksheet):
    """
    Adjusts column width based on the maximum content length in each column.
    This version safely handles merged cells from pandas multi-index headers.
    """
    # Iterate through columns by index (1-based) to avoid issues with merged cells
    for col_idx in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0

        # Find the maximum length of any cell in the current column
        for cell in worksheet[column_letter]:
            try:
                # Check if the cell has a value and find its string length
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                # Ignore any cell that might cause an error
                pass

        # Add a little extra padding to the calculated max length
        adjusted_width = max_length + 3
        worksheet.column_dimensions[column_letter].width = adjusted_width
# ==============================================================================
# ### END OF FIXED FUNCTION ###
# ==============================================================================

def create_excel_report(report_data, sheet_name, df_audit, output_filename):
    """Writes a single-sheet report plus an audit log."""
    try:
        with pd.ExcelWriter(output_filename, engine="openpyxl") as writer:
            # Write the main report sheet
            index_data = [row["Category"] for row in report_data]
            df_data = {
                ("Total Switches", ""): [r["Total Switches"] for r in report_data],
                ("Total Used Ports", ""): [r["Total Used Ports"] for r in report_data],
                ("Total Unused Ports", ""): [r["Total Unused Ports"] for r in report_data],
                ("Proposed No.of Switches", "24 Port"): [r["Proposed 24 Port"] for r in report_data],
                ("Proposed No.of Switches", "48 Port"): [r["Proposed 48 Port"] for r in report_data],
                ("Remarks/Placement of Switches", ""): [r["Remarks"] for r in report_data],
            }
            df = pd.DataFrame(df_data, index=index_data)
            df.index.name = ""
            df.to_excel(writer, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            auto_adjust_column_width(worksheet)

            # Write the scoped Audit Log sheet
            if not df_audit.empty:
                df_audit.to_excel(writer, sheet_name="Device Audit Log", index=False)
                audit_sheet = writer.sheets["Device Audit Log"]
                auto_adjust_column_width(audit_sheet)

        logging.info(f"Successfully generated Excel report: '{output_filename}'")
    except Exception as e:
        logging.error(f"Failed to write Excel file: {e}", exc_info=True)

def main():
    if not os.path.exists(DB_FILENAME):
        logging.error(f"Database file '{DB_FILENAME}' not found.")
        return

    conn = get_db_connection(DB_FILENAME)
    if not conn:
        return

    try:
        cursor = conn.cursor()

        selected_site = _prompt_user_for_choice("Select a Site", fetch_distinct(cursor, "SELECT DISTINCT site FROM lab_areas"))
        if not selected_site: return

        selected_building = _prompt_user_for_choice("Select a Building", fetch_distinct(cursor, "SELECT DISTINCT building FROM lab_areas WHERE site = ?", (selected_site,)))
        if not selected_building: return

        selected_floor = _prompt_user_for_choice("Select a Floor", fetch_distinct(cursor, "SELECT DISTINCT floor FROM lab_areas WHERE site = ? AND building = ?", (selected_site, selected_building)))
        if not selected_floor: return

        lab_areas_on_floor = fetch_distinct(cursor, "SELECT DISTINCT lab_area_name FROM lab_areas WHERE site = ? AND building = ? AND floor = ?", (selected_site, selected_building, selected_floor))
        scope_choices = ["[ALL] - Generate report for the entire floor"] + lab_areas_on_floor
        final_choice = _prompt_user_for_choice("Select a scope for the report", scope_choices)
        if not final_choice: return

        if final_choice.startswith("[ALL]"):
            granularity = "floor"
            area_key = (selected_site, selected_building, selected_floor)
            sheet_name = f"Floor-{'-'.join(area_key)}"
            output_filename = f"Report-Floor-{'-'.join(area_key)}.xlsx"
        else:
            granularity = "lab"
            area_key = (selected_site, selected_building, selected_floor, final_choice)
            sheet_name = f"Lab-{'-'.join(area_key)}"
            output_filename = f"Report-Lab-{'-'.join(area_key)}.xlsx"

        output_filename = re.sub(r'[\\/*?:"<>|]', "-", output_filename)
        sheet_name = sheet_name[:31]

        logging.info(f"Generating report for scope: {area_key} with granularity '{granularity}'")

        queries = get_queries_for_granularity(granularity)
        report_data = generate_report_for_scope(cursor, area_key, queries)

        if not report_data or sum(row["Total Switches"] for row in report_data) == 0:
            logging.warning(f"No qualifying devices found for the selected scope '{area_key}'. No report will be generated.")
            return

        df_audit = generate_scoped_audit_log(cursor, granularity, area_key)

        create_excel_report(report_data, sheet_name, df_audit, output_filename)

    except Exception as e:
        logging.error(f"An unexpected error occurred: {e}", exc_info=True)
    finally:
        if conn:
            conn.close()
            logging.info("Database connection closed.")
        print("\nExiting.")

if __name__ == "__main__":
    main()
