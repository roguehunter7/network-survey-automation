# generate_vendor_summary_report.py
"""
A tool to generate a summary report of device counts by vendor for each lab area.
This script creates a multi-sheet Excel file, with one sheet per site, AND an
additional sheet detailing all devices that were excluded from the counts.
Includes an exception for the hostname 'sw-a1-tmp'.
"""

import logging
import os
import re
import sqlite3

import pandas as pd
from openpyxl.utils import get_column_letter

# --- Configuration ---
DB_FILENAME = "network_survey.db"
EXCEL_FILENAME = "vendor_summary_by_site.xlsx"

# --- VENDOR NORMALIZATION MAPPING ---
VENDOR_MAP = {
    "Cisco": ["cisco"],
    "D-Link": [
        "d-link",
        "dlink",
        "d - link",
        "dllink",
        "d- link",
        "dx-1008",
        "delink",
        "dxs-1800-28p",
    ],
    "Netgear": ["netgear", "netger"],
    "TP-Link": ["tp-link", "tp link", "tp - link", "tplink"],
    "HP / Aruba": ["hp", "aruba", "hewlett packard", "procurve"],
    "Extreme Networks": ["extreme"],
    "Linksys": ["linksys"],
    "Dell": ["dell"],
    "Huawei": ["huawei"],
    "Juniper": ["juniper"],
    "Arista": ["arista"],
    "Ubiquiti": ["ubiquiti", "unifi"],
    "3COM": ["3com"],
    "Nortel": ["nortel"],
    "TrendNet": ["trendnet", "trendnedt"],
    "Digisol": ["digisol"],
    "Asus": ["asus"],
    "WTI": ["wti", "western telematic"],
    "IFS / Interlogix": ["ifs", "interlogix"],
    "Magnum": ["magnum"],
    "Tejas Networks": ["tejas"],
}

# --- Logging Setup ---
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)


def get_db_connection(db_file):
    """Establishes a read-only connection to the SQLite database."""
    if not os.path.exists(db_file):
        logging.error(f"Database file '{db_file}' not found.")
        return None
    try:
        conn = sqlite3.connect(f"file:{db_file}?mode=ro", uri=True)
        logging.info(
            f"Successfully connected to database '{db_file}' in read-only mode."
        )
        return conn
    except sqlite3.Error as e:
        logging.error(f"Error connecting to database '{db_file}': {e}")
        return None


def normalize_vendor_name(vendor_string):
    """Normalizes a vendor name using the VENDOR_MAP."""
    if (
        not vendor_string
        or not isinstance(vendor_string, str)
        or vendor_string.strip().lower() in ["", "unknown"]
    ):
        return "Unknown"

    lower_vendor = vendor_string.lower()
    for canonical_name, variations in VENDOR_MAP.items():
        if any(variation in lower_vendor for variation in variations):
            return canonical_name
    return vendor_string.strip().title()


def sanitize_sheet_name(name):
    """Sanitizes a string to be a valid Excel sheet name."""
    if not name:
        return "Unnamed"
    name = re.sub(r"[\\/*?:\[\]]", "", name)
    return name[:31]


def fetch_all_device_makes(conn):
    """Fetches device data for INCLUDED devices."""
    query = """
    SELECT site, building, floor, lab_area_name, make FROM switches
    WHERE
        site IS NOT NULL AND building IS NOT NULL AND floor IS NOT NULL AND lab_area_name IS NOT NULL
        AND (
            LOWER(COALESCE(hostname, '')) NOT LIKE '%swp%' AND
            LOWER(COALESCE(hostname, '')) NOT LIKE '%rss%' AND
            LOWER(COALESCE(hostname, '')) NOT LIKE '%sws%' OR
            LOWER(COALESCE(hostname, '')) = 'sw-a1-tmp'
        )
        AND NOT (LOWER(COALESCE(make, '')) LIKE '%cisco%' AND (LOWER(COALESCE(model, '')) LIKE '%9200%' OR LOWER(COALESCE(model, '')) LIKE '%9300%' OR LOWER(COALESCE(model, '')) LIKE '%9400%' OR LOWER(COALESCE(model, '')) LIKE '%9500%' OR LOWER(COALESCE(model, '')) LIKE '%9600%'))
    UNION ALL
    SELECT site, building, floor, lab_area_name, reported_make AS make FROM logged_access_points
    WHERE site IS NOT NULL AND building IS NOT NULL AND floor IS NOT NULL AND lab_area_name IS NOT NULL
    UNION ALL
    SELECT site, building, floor, lab_area_name, reported_make AS make FROM logged_other_devices
    WHERE
        site IS NOT NULL AND building IS NOT NULL AND floor IS NOT NULL AND lab_area_name IS NOT NULL
        AND NOT (LOWER(COALESCE(reported_make, '')) LIKE '%cisco%' AND (LOWER(COALESCE(reported_model, '')) LIKE '%9200%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9300%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9400%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9500%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9600%'));
    """
    try:
        df = pd.read_sql_query(query, conn)
        logging.info(f"Fetched {len(df)} INCLUDED device entries for summary.")
        return df
    except Exception as e:
        logging.error(f"Failed to fetch device data for summary: {e}")
        return pd.DataFrame()


def fetch_excluded_devices_details(conn):
    """Fetches the details of all devices that were excluded from the main summary."""
    query = """
    SELECT
        site, building, floor, lab_area_name, hostname, make, model, 'Managed Switch' AS source_table,
        CASE
            WHEN LOWER(COALESCE(hostname, '')) LIKE '%swp%' THEN 'Hostname contains swp'
            WHEN LOWER(COALESCE(hostname, '')) LIKE '%rss%' THEN 'Hostname contains rss'
            WHEN LOWER(COALESCE(hostname, '')) LIKE '%sws%' THEN 'Hostname contains sws'
            WHEN LOWER(COALESCE(make, '')) LIKE '%cisco%' THEN 'Cisco 9k+ Series Model'
            ELSE 'Unknown'
        END AS exclusion_reason
    FROM switches
    WHERE
        site IS NOT NULL AND building IS NOT NULL AND floor IS NOT NULL AND lab_area_name IS NOT NULL AND (
            (
                LOWER(COALESCE(hostname, '')) LIKE '%swp%' OR
                LOWER(COALESCE(hostname, '')) LIKE '%rss%' OR
                LOWER(COALESCE(hostname, '')) LIKE '%sws%'
            ) AND LOWER(COALESCE(hostname, '')) != 'sw-a1-tmp'
            OR (LOWER(COALESCE(make, '')) LIKE '%cisco%' AND (LOWER(COALESCE(model, '')) LIKE '%9200%' OR LOWER(COALESCE(model, '')) LIKE '%9300%' OR LOWER(COALESCE(model, '')) LIKE '%9400%' OR LOWER(COALESCE(model, '')) LIKE '%9500%' OR LOWER(COALESCE(model, '')) LIKE '%9600%'))
        )
    UNION ALL
    SELECT
        site, building, floor, lab_area_name, reported_serial_number as hostname, reported_make as make, reported_model as model, 'Other Device' AS source_table,
        'Cisco 9k+ Series Model' AS exclusion_reason
    FROM logged_other_devices
    WHERE
        site IS NOT NULL AND building IS NOT NULL AND floor IS NOT NULL AND lab_area_name IS NOT NULL
        AND (LOWER(COALESCE(reported_make, '')) LIKE '%cisco%' AND (LOWER(COALESCE(reported_model, '')) LIKE '%9200%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9300%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9400%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9500%' OR LOWER(COALESCE(reported_model, '')) LIKE '%9600%'));
    """
    try:
        df = pd.read_sql_query(query, conn)
        logging.info(f"Fetched details for {len(df)} EXCLUDED devices.")
        return df
    except Exception as e:
        logging.error(f"Failed to fetch excluded device details: {e}")
        return pd.DataFrame()


def auto_adjust_column_width(worksheet):
    """Adjusts column width based on the maximum content length."""
    for col_idx in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        header_cell = worksheet.cell(row=1, column=col_idx)
        if header_cell.value:
            max_length = len(str(header_cell.value))

        for cell in worksheet[column_letter]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = max_length + 3
        worksheet.column_dimensions[column_letter].width = adjusted_width


def main():
    """Main function to drive the summary generation."""
    logging.info("--- Starting Vendor Summary Report Generator ---")

    conn = get_db_connection(DB_FILENAME)
    if not conn:
        return

    try:
        included_devices_df = fetch_all_device_makes(conn)
        excluded_devices_df = fetch_excluded_devices_details(conn)

        if included_devices_df.empty:
            logging.warning(
                "No includable device data found. Summary sheets may be empty."
            )

        included_devices_df["Vendor"] = included_devices_df["make"].apply(
            normalize_vendor_name
        )

        sites = sorted(included_devices_df["site"].unique())

        logging.info(f"Writing report to Excel file: '{EXCEL_FILENAME}'")
        with pd.ExcelWriter(EXCEL_FILENAME, engine="openpyxl") as writer:
            if not included_devices_df.empty:
                for site in sites:
                    sheet_name = sanitize_sheet_name(site)
                    logging.info(f"Creating sheet '{sheet_name}' for site '{site}'...")

                    site_df = included_devices_df[
                        included_devices_df["site"] == site
                    ].copy()

                    site_summary_series = site_df.groupby(
                        ["building", "floor", "lab_area_name", "Vendor"]
                    ).size()

                    if site_summary_series.empty:
                        continue

                    site_summary_df = site_summary_series.unstack(
                        level="Vendor", fill_value=0
                    )
                    site_summary_df["Total Devices"] = site_summary_df.sum(axis=1)
                    site_summary_df.reset_index(inplace=True)

                    site_summary_df.to_excel(writer, sheet_name=sheet_name, index=False)
                    auto_adjust_column_width(writer.sheets[sheet_name])

            if not excluded_devices_df.empty:
                logging.info("Creating sheet for excluded devices...")
                excluded_devices_df.to_excel(
                    writer, sheet_name="Excluded Devices", index=False
                )
                auto_adjust_column_width(writer.sheets["Excluded Devices"])
            else:
                logging.info("No excluded devices found to report.")

        logging.info(f"--- Successfully generated Excel report: '{EXCEL_FILENAME}' ---")
    except Exception as e:
        logging.error(f"An unexpected error occurred: {e}", exc_info=True)
    finally:
        if conn:
            conn.close()
            logging.info("Database connection closed.")


if __name__ == "__main__":
    main()
